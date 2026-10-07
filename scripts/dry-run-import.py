"""
LA CA DA NANG - Offline Dry-Run Import Script (Hardened)
========================================================

Reads the 500-row workbook, applies strict normalization rules from
DB_IMPORT_CONTRACT_PROPOSAL.md, validates all rows against the PostgreSQL
schema constraints, and reports results.

NO DATABASE IS TOUCHED. This is a pure read-only validation pass.

Usage:
    python scripts/dry-run-import.py <path-to-xlsx>

Requirements:
    pip install openpyxl   (already verified available)

Output:
    - Console summary of pass/fail per table
    - JSON report at docs/schema/dry-run-report.json
"""

import sys
import os
import json
import hashlib
import math
import re
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
EXPECTED_HASH = "874e7d6d9f5688af85de347051d4fc67a52f44bbd97a249d1f717056f6496bd3"
EXCEL_EPOCH = datetime(1899, 12, 30)
VALID_SECTIONS = {"EAT", "CAFE", "GO", "STAY"}
VALID_LOCALES = {"vi", "en", "ko"}
VALID_DOMAINS = {"EAT", "CAFE", "GO", "STAY", "COMMON"}

# Core sheet names (only these are imported)
CORE_SHEETS = [
    "administrative_units",
    "places",
    "tags",
    "place_tags",
    "tag_translations",
    "place_translations",
]

# Expected row counts per table
EXPECTED_COUNTS = {
    "administrative_units": 94,
    "tags": 36,
    "places": 500,
    "place_tags": 841,
    "tag_translations": 108,
    "place_translations": 1500,
}

# Expected section counts for places
EXPECTED_SECTIONS = {
    "EAT": 134,
    "CAFE": 145,
    "GO": 94,
    "STAY": 127,
}

# Required headers per sheet
REQUIRED_HEADERS = {
    "administrative_units": ["id", "official_code", "official_name", "unit_type", "active", "created_at", "updated_at"],
    "places": [
        "id", "google_place_id", "name", "section", "primary_type", "phone",
        "website_url", "address", "administrative_unit_id", "latitude", "longitude",
        "google_maps_url", "google_verified_at", "rating", "review_count",
        "photo_count", "business_status", "active", "featured",
        "source_updated_at", "created_at", "updated_at",
    ],
    "tags": ["id", "code", "display_name", "domain", "active"],
    "place_tags": ["place_id", "tag_id"],
    "tag_translations": ["tag_id", "locale", "label"],
    "place_translations": [
        "place_id", "locale", "display_name", "primary_type_label",
        "short_description", "created_at", "updated_at",
    ],
}

# NUMERIC(3,2) max decimal digits: integer part 1 digit, fraction 2 digits
# Max value 9.99 but rating CHECK is 0..5, so max is 5.00
RATING_MAX_DECIMAL_PLACES = 2


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def file_sha256(path: str) -> str:
    """Compute SHA256 of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def excel_serial_to_naive_datetime(serial: float) -> datetime:
    """Convert Excel serial number to naive datetime (no timezone)."""
    return EXCEL_EPOCH + timedelta(days=serial)


def parse_rating(value: Any) -> tuple[Decimal | None, str | None]:
    """Parse rating value compatible with PostgreSQL NUMERIC(3,2).
    Returns (parsed_value, error_or_none).
    NULL stays NULL. Rejects silent precision loss beyond 2 decimal places.
    """
    if value is None:
        return None, None
    # Reject bool masquerading as numeric
    if isinstance(value, bool):
        return None, f"boolean is not a valid rating: {value}"
    try:
        d = Decimal(str(value))
        if not d.is_finite():
            return None, f"non-finite rating: {value}"
        if d < 0 or d > 5:
            return None, f"rating out of range [0,5]: {value}"
        # Check NUMERIC(3,2) compatibility: at most 2 decimal places
        # and total precision at most 3 digits
        sign, digits, exponent = d.as_tuple()
        if exponent < -RATING_MAX_DECIMAL_PLACES:
            # More than 2 decimal places - would be silently truncated
            return None, f"rating has more than {RATING_MAX_DECIMAL_PLACES} decimal places (would lose precision in NUMERIC(3,2)): {value}"
        # Check total precision fits NUMERIC(3,2): max integer digit is 1 (0-5)
        # The integer part can be at most 5, which is 1 digit
        return d, None
    except (InvalidOperation, ValueError):
        return None, f"unparseable rating: {value}"


def parse_nonneg_integer(value: Any, field_name: str) -> tuple[int | None, str | None]:
    """Parse a non-negative integer from string/int/float. NULL stays NULL."""
    if value is None:
        return None, None
    # Reject bool
    if isinstance(value, bool):
        return None, f"boolean is not a valid {field_name}: {value}"
    try:
        d = Decimal(str(value))
        if not d.is_finite():
            return None, f"non-finite {field_name}: {value}"
        if d != int(d):
            return None, f"fractional {field_name}: {value}"
        i = int(d)
        if i < 0:
            return None, f"negative {field_name}: {value}"
        return i, None
    except (InvalidOperation, ValueError):
        return None, f"unparseable {field_name}: {value}"


def parse_positive_integer(value: Any, field_name: str) -> tuple[int | None, str | None]:
    """Parse a positive integer ID."""
    if value is None:
        return None, f"NULL {field_name}"
    # Reject bool
    if isinstance(value, bool):
        return None, f"boolean is not a valid {field_name}: {value}"
    try:
        d = Decimal(str(value))
        if not d.is_finite():
            return None, f"non-finite {field_name}: {value}"
        if d != int(d):
            return None, f"fractional {field_name}: {value}"
        i = int(d)
        if i <= 0:
            return None, f"non-positive {field_name}: {value}"
        if i > 2147483647:
            return None, f"{field_name} exceeds INT4 max: {value}"
        return i, None
    except (InvalidOperation, ValueError):
        return None, f"unparseable {field_name}: {value}"


def require_nonempty_text(value: Any, field_name: str) -> tuple[str | None, str | None]:
    """Require non-empty text after stripping. Rejects whitespace-only."""
    if value is None:
        return None, f"NULL {field_name}"
    # Reject bool masquerading as text
    if isinstance(value, bool):
        return None, f"boolean is not valid text for {field_name}: {value}"
    s = str(value).strip()
    if not s:
        return None, f"empty/whitespace-only {field_name}"
    return s, None


def optional_text(value: Any, field_name: str) -> tuple[str | None, str | None]:
    """Optional text: empty/whitespace-only -> NULL, otherwise trimmed string."""
    if value is None:
        return None, None
    if isinstance(value, bool):
        return None, f"boolean is not valid text for {field_name}: {value}"
    s = str(value).strip()
    if not s:
        return None, None  # empty -> NULL
    return s, None


def validate_boolean(value: Any, field_name: str) -> tuple[bool | None, str | None]:
    """Require actual boolean, not string True/False, not numeric 0/1."""
    if value is None:
        return None, f"NULL {field_name}"
    if isinstance(value, bool):
        return value, None
    return None, f"non-boolean {field_name}: {type(value).__name__}={value}"


def validate_timestamp(value: Any, field_name: str, nullable: bool = False) -> tuple[datetime | None, str | None]:
    """Validate Excel serial timestamp. Reject bool values."""
    if value is None:
        if nullable:
            return None, None
        return None, f"NULL {field_name}"
    # Reject bool as timestamp
    if isinstance(value, bool):
        return None, f"boolean is not a valid timestamp for {field_name}: {value}"
    if not isinstance(value, (int, float)):
        return None, f"non-numeric {field_name}: {type(value).__name__}={value}"
    # Reject non-finite
    if math.isnan(value) or math.isinf(value):
        return None, f"non-finite timestamp {field_name}: {value}"
    try:
        dt = excel_serial_to_naive_datetime(float(value))
        return dt, None
    except (OverflowError, ValueError) as e:
        return None, f"invalid timestamp {field_name}: {value} ({e})"


def validate_finite_float(value: Any, field_name: str, min_val: float, max_val: float) -> tuple[float | None, str | None]:
    """Validate a finite float within a range."""
    if value is None:
        return None, f"NULL {field_name}"
    if isinstance(value, bool):
        return None, f"boolean is not a valid {field_name}: {value}"
    if not isinstance(value, (int, float)):
        return None, f"non-numeric {field_name}: {type(value).__name__}={value}"
    f_val = float(value)
    if math.isnan(f_val) or math.isinf(f_val):
        return None, f"non-finite {field_name}: {value}"
    if f_val < min_val or f_val > max_val:
        return None, f"{field_name} out of range [{min_val},{max_val}]: {value}"
    return f_val, None


def validate_google_maps_url(value: str) -> str | None:
    """Validate Google Maps URL. Returns error message or None."""
    if not value:
        return "empty google_maps_url"
    try:
        parsed = urlparse(value)
    except Exception:
        return f"unparseable URL: {value}"
    # Must be HTTPS (or HTTP)
    if parsed.scheme not in ("https", "http"):
        return f"non-http(s) scheme in google_maps_url: {parsed.scheme}"
    # Must be a Google domain
    host = parsed.hostname or ""
    valid_hosts = (
        "maps.google.com", "www.google.com", "google.com",
        "maps.app.goo.gl", "goo.gl",
        "www.google.com.vn", "google.com.vn",
    )
    if not any(host == h or host.endswith("." + h) for h in valid_hosts):
        return f"non-Google host in google_maps_url: {host}"
    return None


def validate_website_url(value: str) -> str | None:
    """Validate optional website URL. Returns error message or None."""
    if not value:
        return None  # NULL is OK for optional
    try:
        parsed = urlparse(value)
    except Exception:
        return f"unparseable website_url: {value}"
    if parsed.scheme not in ("https", "http"):
        return f"non-http(s) scheme in website_url: {parsed.scheme}"
    if not parsed.hostname:
        return f"missing hostname in website_url: {value}"
    return None


def read_sheet_rows(wb: openpyxl.Workbook, sheet_name: str) -> tuple[list[str], list[dict]]:
    """Read a worksheet into a list of dicts. Returns (headers, rows)."""
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    raw_headers = list(next(rows_iter))
    headers = [str(h).strip() if h is not None else "" for h in raw_headers]
    data = []
    for row_values in rows_iter:
        if all(v is None for v in row_values):
            continue  # skip entirely blank rows
        row_dict = {}
        for i, h in enumerate(headers):
            if h and h != "Column1":  # exclude blank accidental column
                row_dict[h] = row_values[i] if i < len(row_values) else None
        data.append(row_dict)
    return headers, data


def validate_headers(sheet_name: str, actual_headers: list[str]) -> list[str]:
    """Validate that all required headers are present. Returns list of errors."""
    required = REQUIRED_HEADERS.get(sheet_name, [])
    # Filter empty/Column1 headers from actual
    actual_clean = [h for h in actual_headers if h and h != "Column1"]
    errors = []
    for req in required:
        if req not in actual_clean:
            errors.append(f"missing required header '{req}' in sheet '{sheet_name}'")
    return errors


# ---------------------------------------------------------------------------
# Table validators
# ---------------------------------------------------------------------------
def validate_administrative_units(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Validate and normalize administrative_units rows."""
    valid, errors = [], []
    seen_ids = set()
    for i, row in enumerate(rows):
        row_errors = []
        row_num = i + 2  # 1-indexed + header

        pid, err = parse_positive_integer(row.get("id"), "id")
        if err:
            row_errors.append(err)
        elif pid in seen_ids:
            row_errors.append(f"duplicate id: {pid}")
        else:
            seen_ids.add(pid)

        name, err = require_nonempty_text(row.get("official_name"), "official_name")
        if err:
            row_errors.append(err)

        utype, err = require_nonempty_text(row.get("unit_type"), "unit_type")
        if err:
            row_errors.append(err)

        active, err = validate_boolean(row.get("active"), "active")
        if err:
            row_errors.append(err)

        cat, err = validate_timestamp(row.get("created_at"), "created_at")
        if err:
            row_errors.append(err)

        uat, err = validate_timestamp(row.get("updated_at"), "updated_at")
        if err:
            row_errors.append(err)

        if row_errors:
            errors.append({"row": row_num, "errors": row_errors})
        else:
            # official_code: empty -> NULL
            oc_raw = row.get("official_code")
            oc_val = None
            if oc_raw is not None:
                oc_str = str(oc_raw).strip()
                oc_val = oc_str if oc_str else None

            valid.append({
                "id": pid,
                "official_code": oc_val,
                "official_name": name,
                "unit_type": utype,
                "active": active,
                "created_at": cat.isoformat() if cat else None,
                "updated_at": uat.isoformat() if uat else None,
            })

    return valid, errors


def validate_tags(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Validate and normalize tags rows."""
    valid, errors = [], []
    seen_ids = set()
    seen_codes = set()
    for i, row in enumerate(rows):
        row_errors = []
        row_num = i + 2

        pid, err = parse_positive_integer(row.get("id"), "id")
        if err:
            row_errors.append(err)
        elif pid in seen_ids:
            row_errors.append(f"duplicate id: {pid}")
        else:
            seen_ids.add(pid)

        code, err = require_nonempty_text(row.get("code"), "code")
        if err:
            row_errors.append(err)
        else:
            # Enforce uppercase per contract
            if code != code.upper():
                row_errors.append(f"tag code is not uppercase: '{code}'")
            elif code in seen_codes:
                row_errors.append(f"duplicate code: {code}")
            else:
                seen_codes.add(code)

        dname, err = require_nonempty_text(row.get("display_name"), "display_name")
        if err:
            row_errors.append(err)

        domain = str(row.get("domain", "")).strip()
        if domain not in VALID_DOMAINS:
            row_errors.append(f"invalid domain: {domain}")

        active, err = validate_boolean(row.get("active"), "active")
        if err:
            row_errors.append(err)

        if row_errors:
            errors.append({"row": row_num, "errors": row_errors})
        else:
            valid.append({
                "id": pid, "code": code, "display_name": dname,
                "domain": domain, "active": active,
            })

    return valid, errors


def validate_places(rows: list[dict], admin_ids: set[int]) -> tuple[list[dict], list[dict]]:
    """Validate and normalize places rows."""
    valid, errors = [], []
    seen_ids = set()
    seen_google_ids = set()
    for i, row in enumerate(rows):
        row_errors = []
        row_num = i + 2

        pid, err = parse_positive_integer(row.get("id"), "id")
        if err:
            row_errors.append(err)
        elif pid in seen_ids:
            row_errors.append(f"duplicate id: {pid}")
        else:
            seen_ids.add(pid)

        gid, err = require_nonempty_text(row.get("google_place_id"), "google_place_id")
        if err:
            row_errors.append(err)
        elif gid in seen_google_ids:
            row_errors.append(f"duplicate google_place_id: {gid}")
        else:
            seen_google_ids.add(gid)

        name, err = require_nonempty_text(row.get("name"), "name")
        if err:
            row_errors.append(err)

        section = str(row.get("section", "")).strip()
        if section not in VALID_SECTIONS:
            row_errors.append(f"invalid section: {section}")

        ptype, err = require_nonempty_text(row.get("primary_type"), "primary_type")
        if err:
            row_errors.append(err)

        address, err = require_nonempty_text(row.get("address"), "address")
        if err:
            row_errors.append(err)

        admin_id, err = parse_positive_integer(row.get("administrative_unit_id"), "administrative_unit_id")
        if err:
            row_errors.append(err)
        elif admin_id not in admin_ids:
            row_errors.append(f"FK violation: administrative_unit_id={admin_id} not in administrative_units")

        # latitude (finite, -90..90)
        lat, err = validate_finite_float(row.get("latitude"), "latitude", -90, 90)
        if err:
            row_errors.append(err)

        # longitude (finite, -180..180)
        lng, err = validate_finite_float(row.get("longitude"), "longitude", -180, 180)
        if err:
            row_errors.append(err)

        # Google Maps URL
        maps_url, err = require_nonempty_text(row.get("google_maps_url"), "google_maps_url")
        if err:
            row_errors.append(err)
        else:
            maps_err = validate_google_maps_url(maps_url)
            if maps_err:
                row_errors.append(maps_err)

        gv_at, err = validate_timestamp(row.get("google_verified_at"), "google_verified_at", nullable=True)
        if err:
            row_errors.append(err)

        rating, err = parse_rating(row.get("rating"))
        if err:
            row_errors.append(err)

        review_count, err = parse_nonneg_integer(row.get("review_count"), "review_count")
        if err:
            row_errors.append(err)

        photo_count, err = parse_nonneg_integer(row.get("photo_count"), "photo_count")
        if err:
            row_errors.append(err)

        bstatus, err = require_nonempty_text(row.get("business_status"), "business_status")
        if err:
            row_errors.append(err)

        active, err = validate_boolean(row.get("active"), "active")
        if err:
            row_errors.append(err)

        featured, err = validate_boolean(row.get("featured"), "featured")
        if err:
            row_errors.append(err)

        su_at, err = validate_timestamp(row.get("source_updated_at"), "source_updated_at", nullable=True)
        if err:
            row_errors.append(err)

        cat, err = validate_timestamp(row.get("created_at"), "created_at")
        if err:
            row_errors.append(err)

        uat, err = validate_timestamp(row.get("updated_at"), "updated_at")
        if err:
            row_errors.append(err)

        # phone: optional text, empty -> NULL
        phone_val, _ = optional_text(row.get("phone"), "phone")

        # website_url: optional text with URL validation
        website_raw, _ = optional_text(row.get("website_url"), "website_url")
        if website_raw:
            ws_err = validate_website_url(website_raw)
            if ws_err:
                row_errors.append(ws_err)

        if row_errors:
            errors.append({"row": row_num, "errors": row_errors})
        else:
            valid.append({
                "id": pid,
                "google_place_id": gid,
                "name": name,
                "section": section,
                "primary_type": ptype,
                "phone": phone_val,
                "website_url": website_raw,
                "address": address,
                "administrative_unit_id": admin_id,
                "latitude": lat,
                "longitude": lng,
                "google_maps_url": maps_url,
                "google_verified_at": gv_at.isoformat() if gv_at else None,
                "rating": str(rating) if rating is not None else None,
                "review_count": review_count,
                "photo_count": photo_count,
                "business_status": bstatus,
                "active": active,
                "featured": featured,
                "source_updated_at": su_at.isoformat() if su_at else None,
                "created_at": cat.isoformat() if cat else None,
                "updated_at": uat.isoformat() if uat else None,
            })

    return valid, errors


def validate_place_tags(rows: list[dict], place_ids: set[int], tag_ids: set[int]) -> tuple[list[dict], list[dict]]:
    """Validate place_tags junction rows."""
    valid, errors = [], []
    seen_pks = set()
    for i, row in enumerate(rows):
        row_errors = []
        row_num = i + 2

        place_id, err = parse_positive_integer(row.get("place_id"), "place_id")
        if err:
            row_errors.append(err)
        elif place_id not in place_ids:
            row_errors.append(f"FK violation: place_id={place_id} not in places")

        tag_id, err = parse_positive_integer(row.get("tag_id"), "tag_id")
        if err:
            row_errors.append(err)
        elif tag_id not in tag_ids:
            row_errors.append(f"FK violation: tag_id={tag_id} not in tags")

        if not row_errors:
            pk = (place_id, tag_id)
            if pk in seen_pks:
                row_errors.append(f"duplicate PK: ({place_id},{tag_id})")
            else:
                seen_pks.add(pk)

        if row_errors:
            errors.append({"row": row_num, "errors": row_errors})
        else:
            valid.append({"place_id": place_id, "tag_id": tag_id})

    return valid, errors


def validate_tag_translations(rows: list[dict], tag_ids: set[int]) -> tuple[list[dict], list[dict]]:
    """Validate tag_translations rows."""
    valid, errors = [], []
    seen_pks = set()
    for i, row in enumerate(rows):
        row_errors = []
        row_num = i + 2

        tag_id, err = parse_positive_integer(row.get("tag_id"), "tag_id")
        if err:
            row_errors.append(err)
        elif tag_id not in tag_ids:
            row_errors.append(f"FK violation: tag_id={tag_id} not in tags")

        locale = str(row.get("locale", "")).strip()
        if locale not in VALID_LOCALES:
            row_errors.append(f"invalid locale: {locale}")

        label, err = require_nonempty_text(row.get("label"), "label")
        if err:
            row_errors.append(err)

        if not row_errors:
            pk = (tag_id, locale)
            if pk in seen_pks:
                row_errors.append(f"duplicate PK: ({tag_id},{locale})")
            else:
                seen_pks.add(pk)

        if row_errors:
            errors.append({"row": row_num, "errors": row_errors})
        else:
            valid.append({"tag_id": tag_id, "locale": locale, "label": label})

    return valid, errors


def validate_place_translations(rows: list[dict], place_ids: set[int]) -> tuple[list[dict], list[dict]]:
    """Validate place_translations rows."""
    valid, errors = [], []
    seen_pks = set()
    for i, row in enumerate(rows):
        row_errors = []
        row_num = i + 2

        place_id, err = parse_positive_integer(row.get("place_id"), "place_id")
        if err:
            row_errors.append(err)
        elif place_id not in place_ids:
            row_errors.append(f"FK violation: place_id={place_id} not in places")

        locale = str(row.get("locale", "")).strip()
        if locale not in VALID_LOCALES:
            row_errors.append(f"invalid locale: {locale}")

        dname, err = require_nonempty_text(row.get("display_name"), "display_name")
        if err:
            row_errors.append(err)

        ptlabel, err = require_nonempty_text(row.get("primary_type_label"), "primary_type_label")
        if err:
            row_errors.append(err)

        # Validate actual short_description cell instead of hard-coding None
        sd_val, sd_err = optional_text(row.get("short_description"), "short_description")
        if sd_err:
            row_errors.append(sd_err)

        cat, err = validate_timestamp(row.get("created_at"), "created_at")
        if err:
            row_errors.append(err)

        uat, err = validate_timestamp(row.get("updated_at"), "updated_at")
        if err:
            row_errors.append(err)

        if not row_errors:
            pk = (place_id, locale)
            if pk in seen_pks:
                row_errors.append(f"duplicate PK: ({place_id},{locale})")
            else:
                seen_pks.add(pk)

        if row_errors:
            errors.append({"row": row_num, "errors": row_errors})
        else:
            valid.append({
                "place_id": place_id,
                "locale": locale,
                "display_name": dname,
                "primary_type_label": ptlabel,
                "short_description": sd_val,  # actual cell value, NULL if empty
                "created_at": cat.isoformat() if cat else None,
                "updated_at": uat.isoformat() if uat else None,
            })

    return valid, errors


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/dry-run-import.py <path-to-xlsx>")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    if not os.path.isfile(xlsx_path):
        print(f"ERROR: File not found: {xlsx_path}")
        sys.exit(1)

    # ===================================================================
    # 1. Hash verification
    # ===================================================================
    print("=" * 60)
    print("LA CA DA NANG - Dry-Run Import Validation (Hardened)")
    print("=" * 60)
    print(f"\nFile: {xlsx_path}")

    actual_hash = file_sha256(xlsx_path)
    hash_match = actual_hash == EXPECTED_HASH
    print(f"SHA256: {actual_hash}")
    print(f"Expected: {EXPECTED_HASH}")
    print(f"Hash match: {'PASS' if hash_match else 'FAIL'}")
    if not hash_match:
        print("WARNING: Hash mismatch! This may not be the expected workbook.")

    report = {
        "timestamp": datetime.now().isoformat(),
        "file": xlsx_path,
        "hash_check": {
            "actual": actual_hash,
            "expected": EXPECTED_HASH,
            "match": hash_match,
        },
    }

    # ===================================================================
    # 2. Open workbook and validate sheets/headers
    # ===================================================================
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    print(f"\nSheets found: {sheet_names}")

    missing_sheets = [s for s in CORE_SHEETS if s not in sheet_names]
    if missing_sheets:
        print(f"ERROR: Missing required sheets: {missing_sheets}")
        wb.close()
        sys.exit(1)

    # Validate headers for all core sheets
    header_checks = {}
    all_header_errors = []
    sheet_data = {}
    for sheet_name in CORE_SHEETS:
        headers, rows = read_sheet_rows(wb, sheet_name)
        sheet_data[sheet_name] = (headers, rows)
        h_errors = validate_headers(sheet_name, headers)
        header_checks[sheet_name] = {
            "actual_headers": [h for h in headers if h and h != "Column1"],
            "required_headers": REQUIRED_HEADERS[sheet_name],
            "errors": h_errors,
        }
        if h_errors:
            all_header_errors.extend(h_errors)
            for e in h_errors:
                print(f"  HEADER ERROR: {e}")

    report["schema_header_checks"] = header_checks
    print(f"\nHeader validation: {'PASS' if not all_header_errors else 'FAIL (' + str(len(all_header_errors)) + ' errors)'}")

    # ===================================================================
    # 3. Validate each table in dependency order
    # ===================================================================
    print("\n" + "-" * 60)
    print("Validating administrative_units...")
    _, admin_rows = sheet_data["administrative_units"]
    admin_valid, admin_errors = validate_administrative_units(admin_rows)
    admin_ids = {r["id"] for r in admin_valid}
    print(f"  Total: {len(admin_rows)} | Valid: {len(admin_valid)} | Errors: {len(admin_errors)}")

    print("\nValidating tags...")
    _, tag_rows = sheet_data["tags"]
    tag_valid, tag_errors = validate_tags(tag_rows)
    tag_ids = {r["id"] for r in tag_valid}
    print(f"  Total: {len(tag_rows)} | Valid: {len(tag_valid)} | Errors: {len(tag_errors)}")

    print("\nValidating places...")
    _, place_rows = sheet_data["places"]
    place_valid, place_errors = validate_places(place_rows, admin_ids)
    place_ids = {r["id"] for r in place_valid}
    section_counts = {}
    for p in place_valid:
        section_counts[p["section"]] = section_counts.get(p["section"], 0) + 1
    print(f"  Total: {len(place_rows)} | Valid: {len(place_valid)} | Errors: {len(place_errors)}")
    print(f"  Sections: {section_counts}")

    print("\nValidating place_tags...")
    _, pt_rows = sheet_data["place_tags"]
    pt_valid, pt_errors = validate_place_tags(pt_rows, place_ids, tag_ids)
    print(f"  Total: {len(pt_rows)} | Valid: {len(pt_valid)} | Errors: {len(pt_errors)}")

    print("\nValidating tag_translations...")
    _, tt_rows = sheet_data["tag_translations"]
    tt_valid, tt_errors = validate_tag_translations(tt_rows, tag_ids)
    print(f"  Total: {len(tt_rows)} | Valid: {len(tt_valid)} | Errors: {len(tt_errors)}")

    print("\nValidating place_translations...")
    _, ptr_rows = sheet_data["place_translations"]
    ptr_valid, ptr_errors = validate_place_translations(ptr_rows, place_ids)
    print(f"  Total: {len(ptr_rows)} | Valid: {len(ptr_valid)} | Errors: {len(ptr_errors)}")

    wb.close()

    # Build table validation report
    table_validation = {
        "administrative_units": {"total": len(admin_rows), "valid": len(admin_valid), "errors": admin_errors},
        "tags": {"total": len(tag_rows), "valid": len(tag_valid), "errors": tag_errors},
        "places": {"total": len(place_rows), "valid": len(place_valid), "sections": section_counts, "errors": place_errors},
        "place_tags": {"total": len(pt_rows), "valid": len(pt_valid), "errors": pt_errors},
        "tag_translations": {"total": len(tt_rows), "valid": len(tt_valid), "errors": tt_errors},
        "place_translations": {"total": len(ptr_rows), "valid": len(ptr_valid), "errors": ptr_errors},
    }
    report["table_validation"] = table_validation

    # ===================================================================
    # 4. Row count checks
    # ===================================================================
    print("\n" + "-" * 60)
    print("Row count checks...")
    actual_counts = {
        "administrative_units": len(admin_rows),
        "tags": len(tag_rows),
        "places": len(place_rows),
        "place_tags": len(pt_rows),
        "tag_translations": len(tt_rows),
        "place_translations": len(ptr_rows),
    }
    count_mismatches = []
    for table, expected in EXPECTED_COUNTS.items():
        actual = actual_counts.get(table, 0)
        match = actual == expected
        if not match:
            count_mismatches.append({"table": table, "expected": expected, "actual": actual})
        print(f"  {table}: {actual}/{expected} {'PASS' if match else 'FAIL'}")

    report["row_count_checks"] = {
        "expected": EXPECTED_COUNTS,
        "actual": actual_counts,
        "mismatches": count_mismatches,
        "all_match": len(count_mismatches) == 0,
    }

    # ===================================================================
    # 5. Section count checks
    # ===================================================================
    print("\nSection count checks...")
    section_mismatches = []
    for section, expected in EXPECTED_SECTIONS.items():
        actual = section_counts.get(section, 0)
        match = actual == expected
        if not match:
            section_mismatches.append({"section": section, "expected": expected, "actual": actual})
        print(f"  {section}: {actual}/{expected} {'PASS' if match else 'FAIL'}")

    report["section_count_checks"] = {
        "expected": EXPECTED_SECTIONS,
        "actual": section_counts,
        "mismatches": section_mismatches,
        "all_match": len(section_mismatches) == 0,
    }

    # ===================================================================
    # 6. Cross-table checks
    # ===================================================================
    print("\n" + "-" * 60)
    print("Cross-table referential checks...")

    # Place locale completeness
    place_locales = {}
    for ptr in ptr_valid:
        pid = ptr["place_id"]
        if pid not in place_locales:
            place_locales[pid] = set()
        place_locales[pid].add(ptr["locale"])

    places_missing_locales = []
    for pid in sorted(place_ids):
        locales = place_locales.get(pid, set())
        missing_loc = VALID_LOCALES - locales
        if missing_loc:
            places_missing_locales.append({"place_id": pid, "missing": sorted(missing_loc)})

    print(f"  Places missing locale translations: {len(places_missing_locales)}")

    # Tag locale completeness
    tag_locales = {}
    for tt in tt_valid:
        tid = tt["tag_id"]
        if tid not in tag_locales:
            tag_locales[tid] = set()
        tag_locales[tid].add(tt["locale"])

    tags_missing_locales = []
    for tid in sorted(tag_ids):
        locales = tag_locales.get(tid, set())
        missing_loc = VALID_LOCALES - locales
        if missing_loc:
            tags_missing_locales.append({"tag_id": tid, "missing": sorted(missing_loc)})

    print(f"  Tags missing locale translations: {len(tags_missing_locales)}")

    # Places without tags
    tagged_place_ids = {pt["place_id"] for pt in pt_valid}
    untagged = sorted(place_ids - tagged_place_ids)
    print(f"  Places without any tags: {len(untagged)}")

    cross_checks = {
        "places_missing_locale_translations": places_missing_locales,
        "tags_missing_locale_translations": tags_missing_locales,
        "untagged_places": untagged,
        "places_missing_locale_count": len(places_missing_locales),
        "tags_missing_locale_count": len(tags_missing_locales),
        "untagged_place_count": len(untagged),
    }
    report["cross_checks"] = cross_checks

    # ===================================================================
    # 7. Summary — strict all_pass gate
    # ===================================================================
    print("\n" + "=" * 60)
    total_errors = sum(len(t["errors"]) for t in table_validation.values())
    total_rows = sum(t["total"] for t in table_validation.values())
    total_valid = sum(t["valid"] for t in table_validation.values())

    # Strict all_pass: ALL conditions must be true
    all_pass = (
        hash_match
        and len(all_header_errors) == 0
        and total_errors == 0
        and len(count_mismatches) == 0
        and len(section_mismatches) == 0
        and len(places_missing_locales) == 0
        and len(tags_missing_locales) == 0
        and len(untagged) == 0
    )

    report["summary"] = {
        "total_rows": total_rows,
        "total_valid": total_valid,
        "total_errors": total_errors,
        "hash_pass": hash_match,
        "headers_pass": len(all_header_errors) == 0,
        "counts_pass": len(count_mismatches) == 0,
        "sections_pass": len(section_mismatches) == 0,
        "locale_completeness_pass": len(places_missing_locales) == 0 and len(tags_missing_locales) == 0,
        "all_tagged_pass": len(untagged) == 0,
        "all_pass": all_pass,
    }

    if all_pass:
        print(f"RESULT: ALL PASS - {total_valid}/{total_rows} rows validated")
    else:
        print(f"RESULT: FAIL")
        if not hash_match:
            print("  - Hash mismatch")
        if all_header_errors:
            print(f"  - {len(all_header_errors)} header errors")
        if total_errors > 0:
            print(f"  - {total_errors} row validation errors")
        if count_mismatches:
            print(f"  - {len(count_mismatches)} row count mismatches")
        if section_mismatches:
            print(f"  - {len(section_mismatches)} section count mismatches")
        if places_missing_locales:
            print(f"  - {len(places_missing_locales)} places missing locale translations")
        if tags_missing_locales:
            print(f"  - {len(tags_missing_locales)} tags missing locale translations")
        if untagged:
            print(f"  - {len(untagged)} untagged places")

    # ===================================================================
    # 8. Write report
    # ===================================================================
    report_dir = Path(__file__).resolve().parent.parent / "docs" / "schema"
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / "dry-run-report.json"
    with open(report_path, "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2, ensure_ascii=False, default=str)
    print(f"\nReport written to: {report_path}")
    print("=" * 60)

    sys.exit(0 if all_pass else 1)


if __name__ == "__main__":
    main()
